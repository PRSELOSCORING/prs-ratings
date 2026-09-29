"""
PRS Elo rating engine.

Ported from Clay's Buschman_Rating_2026.xlsx and validated bit-for-bit against
five real matches (Best In Texas 26, Okie 26, Vortex Vengeance 26, Magnolia 26,
Frostbite 26) -- 647 shooter-results, max difference 0.0000000000.

Requires: openpyxl (only for the validation helper at the bottom)
"""

SPREAD = 133          # Ratings!B2 -- lower than chess Elo's 400, so ratings move faster
VOLATILITY = 2        # Ratings!B3 -- K-factor, points per pairwise comparison
NEW_GUY_RATING = 500  # Ratings!B4 -- default for first-time competitors
MAX_GAIN = 50         # per-match ceiling
MAX_LOSS = 25         # per-match floor


def compute_match_ratings(
    results,
    spread=SPREAD,
    volatility=VOLATILITY,
    max_gain=MAX_GAIN,
    max_loss=MAX_LOSS,
):
    """
    Compute post-match ratings for every shooter in a single match.

    results: iterable of (place, name, rating_in) tuples.
             `place` is the finishing position, 1 = winner.
             `rating_in` is the shooter's rating entering this match
             (use NEW_GUY_RATING for first-timers).

    Returns: dict mapping name -> new rating (float).

    Algorithm: standard Elo run as a full round-robin over the field.
    Every shooter is compared against every other shooter; beating someone
    you were expected to beat earns almost nothing, beating someone rated
    far above you earns a lot. The total is clamped to [-25, +50].
    """
    field = sorted(results, key=lambda r: r[0])
    n = len(field)
    new_ratings = {}

    for i in range(n):
        place_a, name_a, rating_a = field[i]
        delta = 0.0

        for j in range(n):
            if i == j:
                continue
            place_b, _name_b, rating_b = field[j]

            actual = 1 if place_a < place_b else 0
            expected = 1.0 / (1.0 + 10 ** ((rating_b - rating_a) / spread))
            delta += volatility * (actual - expected)

        raw = rating_a + delta
        new_ratings[name_a] = min(rating_a + max_gain, max(rating_a - max_loss, raw))

    return new_ratings


def rate_one(
    place,
    rating,
    opponents,
    spread=SPREAD,
    volatility=VOLATILITY,
    max_gain=MAX_GAIN,
    max_loss=MAX_LOSS,
):
    """
    New rating for ONE shooter, holding everyone else's ratings fixed.

    opponents: iterable of (place, rating_in) for every other shooter in the match.
    Same math as compute_match_ratings; used for what-ifs such as a mulligan.
    """
    delta = 0.0
    for place_b, rating_b in opponents:
        actual = 1 if place < place_b else 0
        expected = 1.0 / (1.0 + 10 ** ((rating_b - rating) / spread))
        delta += volatility * (actual - expected)
    return min(rating + max_gain, max(rating - max_loss, rating + delta))


def win_probability(rating_a, rating_b, spread=SPREAD):
    """Probability that shooter A beats shooter B. Handy for a head-to-head predictor."""
    return 1.0 / (1.0 + 10 ** ((rating_b - rating_a) / spread))


def run_season(matches, starting_ratings=None, new_guy=NEW_GUY_RATING):
    """
    Run a full season in order.

    matches: list of (match_name, [(place, name), ...]) in CHRONOLOGICAL order.
             Order matters -- each match's inputs are the previous match's outputs.
    starting_ratings: dict name -> rating to seed from (e.g. the 2025 Finale column).

    Returns: (final_ratings dict, history dict keyed by match name)
    """
    ratings = dict(starting_ratings or {})
    history = {}

    for match_name, placings in matches:
        with_ratings = [
            (place, name, ratings.get(name, new_guy)) for place, name in placings
        ]
        updated = compute_match_ratings(with_ratings)

        history[match_name] = [
            {
                "place": place,
                "name": name,
                "rating_in": rating_in,
                "rating_out": updated[name],
                "delta": updated[name] - rating_in,
            }
            for place, name, rating_in in with_ratings
        ]

        ratings.update(updated)

    return ratings, history


# ---------------------------------------------------------------------------
# Parsing Impact Scoring results
# ---------------------------------------------------------------------------

def parse_live_scores(payload):
    """
    Turn an Impact Scoring getLiveScores payload into [(place, name), ...].

    Endpoint:
        https://app-api.impactscoring.net/rest/website/getLiveScores
            ?matchId=<N>&dividePoints=false
    Requires a logged-in session (HttpOnly cookies).

    WARNING: the `dropped` field on each shooter is a POINTS DIFFERENTIAL
    (e.g. -300), not a boolean drop flag. Do not filter on it -- doing so
    discards ~98% of the field. Filter on points > 0 instead.
    """
    shooters = payload.get("shooters", [])
    scored = [s for s in shooters if (s.get("points") or 0) > 0]
    scored.sort(key=lambda s: s.get("rank") or 10**9)

    return [
        (s["rank"], f"{(s.get('firstName') or '').strip()} {(s.get('lastName') or '').strip()}".strip())
        for s in scored
    ]


def normalize_name(name):
    """
    Collapse casing and whitespace so spreadsheet names match Impact Scoring names.

    Real examples that break naive matching:
        "Morgun king"  vs "Morgun King"
        "gage caples"  vs "Gage Caples"
        "Kristy  Caples" (double space)

    This is a starting point, not a complete solution -- nicknames and
    hyphenated surnames will still need a fuzzy-match layer.
    """
    return " ".join((name or "").split()).lower()


# ---------------------------------------------------------------------------
# Validation helper -- proves the engine matches the spreadsheet
# ---------------------------------------------------------------------------

def validate_against_workbook(path, sheet_names, tolerance=0.0001):
    """
    Re-run the validation. Reads cached values from the workbook, so the file
    must have been recalculated and saved by Excel first.

    Match sheet layout: headers on row 2, data from row 3.
        D = Place, E = Name, F = Rating in, G = New Rating
    """
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True)
    report = []

    for sheet_name in sheet_names:
        ws = wb[sheet_name]
        results, expected = [], {}

        for row in range(3, ws.max_row + 1):
            place = ws.cell(row=row, column=4).value
            name = ws.cell(row=row, column=5).value
            rating_in = ws.cell(row=row, column=6).value
            rating_out = ws.cell(row=row, column=7).value
            if None in (place, name, rating_in, rating_out):
                continue
            results.append((place, name, rating_in))
            expected[name] = rating_out

        if not results:
            report.append((sheet_name, 0, None, 0))
            continue

        computed = compute_match_ratings(results)
        diffs = [abs(computed[n] - e) for n, e in expected.items()]
        report.append((
            sheet_name,
            len(computed),
            max(diffs),
            sum(1 for d in diffs if d > tolerance),
        ))

    return report


if __name__ == "__main__":
    import sys

    if len(sys.argv) > 1:
        sheets = [
            "Best In Texas 26",
            "Okie 26",
            "Vortex Vengeance 26",
            "Magnolia 26",
            "Frostbite 26",
        ]
        for name, count, max_diff, bad in validate_against_workbook(sys.argv[1], sheets):
            if max_diff is None:
                print(f"{name}: no cached data (recalculate in Excel first)")
            else:
                print(f"{name}: {count} shooters, max diff {max_diff:.10f}, mismatches {bad}")
    else:
        print(__doc__)
        print("Usage: python prs_rating_engine.py <path-to-workbook.xlsx>")
